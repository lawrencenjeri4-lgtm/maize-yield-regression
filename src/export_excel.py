"""
Build a styled Excel results workbook (outputs/maize_results.xlsx) from the project data.

Usage (from the repo root, after the merge steps have been run):
    python src/export_excel.py

Inputs:
    data/raw/yield_data_county.csv          county, year, yield_bags_per_ha
    data/raw/yield_data_national.csv        county ("Kenya"), year, yield_bags_per_ha
    data/processed/maize_data_county.csv    made by: python src/merge_data.py <county file> <output>
    data/processed/maize_data_national.csv  made the same way for the national file

Output:
    outputs/maize_results.xlsx   (4 sheets: County Ranking, Model Comparison, County Data, National Data)

The regression numbers are calculated here (numpy + scipy), so the workbook is rebuilt
from your data each time you run this. Nothing is typed in by hand.

Install: pip install openpyxl numpy pandas scipy
Open the file in Excel (or Google Sheets): the ranking columns are formulas and
calculate when the file is opened.
"""
import statistics
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy import stats

COUNTY_RAW = "data/raw/yield_data_county.csv"
NATIONAL_RAW = "data/raw/yield_data_national.csv"
COUNTY_MERGED = "data/processed/maize_data_county.csv"
NATIONAL_MERGED = "data/processed/maize_data_national.csv"
OUT_FILE = "outputs/maize_results.xlsx"

# ---------------------------------------------------------------- regression

def ols(X, y):
    """Ordinary least squares with a constant already in X. Returns a dict of results."""
    X, y = np.asarray(X, float), np.asarray(y, float)
    n, k = X.shape
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)
    resid = y - X @ beta
    ssr = float(resid @ resid)
    df = n - k
    cov = (ssr / df) * np.linalg.inv(X.T @ X)
    se = np.sqrt(np.diag(cov))
    t = beta / se
    p = 2 * stats.t.sf(np.abs(t), df)
    tss = float(((y - y.mean()) ** 2).sum())
    r2 = 1 - ssr / tss
    adj = 1 - (1 - r2) * (n - 1) / df
    llf = -n / 2 * (np.log(2 * np.pi) + np.log(ssr / n) + 1)
    aic = -2 * llf + 2 * k
    return {"beta": beta, "se": se, "p": p, "r2": r2, "adj": adj, "aic": aic, "n": n}


def fit_rainfall(df, county_effects=False):
    """Fit yield_kg_per_acre ~ rainfall_mm (+ county dummies). Returns rainfall row stats."""
    df = df.dropna(subset=["yield_kg_per_acre", "rainfall_mm"])
    cols = [np.ones(len(df)), df["rainfall_mm"].to_numpy()]
    if county_effects:
        dummies = pd.get_dummies(df["county"], drop_first=True).astype(float)
        cols += [dummies[c].to_numpy() for c in dummies.columns]
    res = ols(np.column_stack(cols), df["yield_kg_per_acre"].to_numpy())
    return {
        "n": res["n"], "coef": res["beta"][1], "se": res["se"][1], "p": res["p"][1],
        "r2": res["r2"], "adj": res["adj"], "aic": res["aic"],
    }

# ---------------------------------------------------------------- styling

FONT = "Arial"
GREEN_DARK, GREEN_BAND, GREY = "1B5E20", "E8F5E9", "BDBDBD"
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", start_color=GREEN_DARK)
band_fill = PatternFill("solid", start_color=GREEN_BAND)
body_font = Font(name=FONT, size=10)
bold_font = Font(name=FONT, size=10, bold=True)
title_font = Font(name=FONT, size=14, bold=True, color=GREEN_DARK)
sub_font = Font(name=FONT, size=10, italic=True, color="555555")
note_font = Font(name=FONT, size=9, color="444444")
thin = Side(style="thin", color=GREY)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.alignment, cell.border = hdr_font, hdr_fill, center, border
    ws.row_dimensions[row].height = 34


def style_body(ws, first, last, ncols, fmts=None):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border, cell.alignment = body_font, border, center
            if (r - first) % 2 == 1:
                cell.fill = band_fill
            if fmts and c in fmts:
                cell.number_format = fmts[c]


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_notes(ws, start_row, lines):
    for i, text in enumerate(lines):
        c = ws.cell(row=start_row + i, column=1, value=text)
        c.font = bold_font if i == 0 else note_font

# ---------------------------------------------------------------- build

def main():
    county = pd.read_csv(COUNTY_RAW)
    national = pd.read_csv(NATIONAL_RAW)
    m_county = pd.read_csv(COUNTY_MERGED)
    m_national = pd.read_csv(NATIONAL_MERGED)

    wb = Workbook()

    # ---- County Data
    cd = wb.active
    cd.title = "County Data"
    cd.append(["County", "Year", "Yield (90 kg bags per hectare)"])
    for r in county[["county", "year", "yield_bags_per_ha"]].itertuples(index=False):
        cd.append([r[0], int(r[1]), float(r[2])])
    n_cd = len(county) + 1
    style_header(cd, 1, 3)
    style_body(cd, 2, n_cd, 3, {2: "0", 3: "0.0"})
    cd.cell(row=n_cd + 2, column=1, value=(
        "Source: KNBS National Agriculture Production Report (maize area and production by county). "
        "Yield = tonnes / 0.09 / hectares. The latest year may be provisional.")).font = note_font
    widths(cd, [18, 10, 30])
    cd.freeze_panes = "A2"
    cd.auto_filter.ref = f"A1:C{n_cd}"

    # ---- National Data
    nd = wb.create_sheet("National Data")
    nd.append(["Region", "Year", "Yield (90 kg bags per hectare)"])
    for r in national[["county", "year", "yield_bags_per_ha"]].itertuples(index=False):
        nd.append([r[0], int(r[1]), float(r[2])])
    n_nd = len(national) + 1
    style_header(nd, 1, 3)
    style_body(nd, 2, n_nd, 3, {2: "0", 3: "0.0"})
    nd.cell(row=n_nd + 2, column=1, value=(
        "Sources: Ministry of Agriculture Economic Review of Agriculture and KNBS national area and "
        "production. The sources may use different methods, and some years may be missing.")).font = note_font
    widths(nd, [18, 10, 30])
    nd.freeze_panes = "A2"

    # ---- County Ranking
    cr = wb.create_sheet("County Ranking", 0)
    cr["A1"] = "Maize yield by county vs national average"
    cr["A1"].font = title_font
    cr["A2"] = "Average yield in 90 kg bags per hectare, over the years covered by the county data."
    cr["A2"].font = sub_font
    stat_rows = [
        ("First year", f"=MIN('County Data'!$B$2:$B${n_cd})", "0"),
        ("Last year", f"=MAX('County Data'!$B$2:$B${n_cd})", "0"),
        ("National average, same years",
         f"=AVERAGEIFS('National Data'!$C$2:$C${n_nd},'National Data'!$B$2:$B${n_nd},\">=\"&B4,"
         f"'National Data'!$B$2:$B${n_nd},\"<=\"&B5)", "0.0"),
    ]
    for i, (lab, formula, fmt) in enumerate(stat_rows):
        r = 4 + i
        cr.cell(row=r, column=1, value=lab).font = bold_font
        cr.cell(row=r, column=1).alignment = left
        cr.cell(row=r, column=1).border = border
        c = cr.cell(row=r, column=2, value=formula)
        c.font, c.number_format, c.alignment, c.border = body_font, fmt, center, border

    hdr_row = 8
    for j, h in enumerate(["Rank", "County", "Average yield", "Lowest year", "Highest year",
                           "vs national average"], start=1):
        cr.cell(row=hdr_row, column=j, value=h)
    style_header(cr, hdr_row, 6)

    order = sorted(county["county"].unique(),
                   key=lambda c: -statistics.mean(county.loc[county["county"] == c, "yield_bags_per_ha"]))
    first, last = hdr_row + 1, hdr_row + len(order)
    rng = lambda col: f"'County Data'!${col}$2:${col}${n_cd}"
    for i, name in enumerate(order):
        r = first + i
        cr.cell(row=r, column=1, value=f"=RANK(C{r},$C${first}:$C${last})")
        cr.cell(row=r, column=2, value=name)
        cr.cell(row=r, column=3, value=f"=AVERAGEIFS({rng('C')},{rng('A')},B{r})")
        cr.cell(row=r, column=4, value=f"=_xlfn.MINIFS({rng('C')},{rng('A')},B{r})")
        cr.cell(row=r, column=5, value=f"=_xlfn.MAXIFS({rng('C')},{rng('A')},B{r})")
        cr.cell(row=r, column=6, value=f"=C{r}/$B$6-1")
    style_body(cr, first, last, 6, {1: "0", 3: "0.0", 4: "0.0", 5: "0.0", 6: "0%"})
    for r in range(first, last + 1):
        cr.cell(row=r, column=2).font = bold_font
    cr.conditional_formatting.add(
        f"C{first}:C{last}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="43A047", showValue=True))
    add_notes(cr, last + 2, [
        "Notes",
        "These are the counties chosen for this analysis, so this is not a ranking of all counties.",
        "Counties with close averages should be treated as roughly tied.",
        "The latest year in the data may be provisional. See the County Data sheet for sources.",
    ])
    widths(cr, [28, 16, 16, 14, 14, 20])
    cr.freeze_panes = "A9"

    # ---- Model Comparison
    m1 = fit_rainfall(m_national)
    m2 = fit_rainfall(m_county)
    m3 = fit_rainfall(m_county, county_effects=True)
    ny = m_national["year"].nunique()
    nc = m_county["county"].nunique()
    yrs = f"{int(m_county['year'].min())}-{int(m_county['year'].max())}"
    models = [
        ("1. Rainfall only", f"National, {ny} years", m1),
        ("2. Rainfall only", f"{nc} counties, {yrs}", m2),
        ("3. Rainfall + county effects", f"{nc} counties, {yrs}", m3),
    ]

    mc = wb.create_sheet("Model Comparison", 1)
    mc["A1"] = "Maize yield vs rainfall: model comparison"
    mc["A1"].font = title_font
    mc["A2"] = "Outcome: maize yield (kg per acre). Predictor: seasonal rainfall (mm). Calculated by src/export_excel.py."
    mc["A2"].font = sub_font
    heads = ["Model", "Data", "Observations (n)", "Rainfall coefficient (kg/acre per mm)", "Std. error",
             "p-value", "Effect of +100 mm rain (kg/acre)", "R\u00b2", "Adjusted R\u00b2", "AIC",
             "Significant at 5%?"]
    hr = 4
    for j, h in enumerate(heads, start=1):
        mc.cell(row=hr, column=j, value=h)
    style_header(mc, hr, len(heads))
    for i, (name, data, m) in enumerate(models):
        r = hr + 1 + i
        values = [name, data, int(m["n"]), float(m["coef"]), float(m["se"]), float(m["p"]),
                  f"=D{r}*100", float(m["r2"]), float(m["adj"]), float(m["aic"]),
                  f'=IF(F{r}<0.05,"Yes","No")']
        for j, v in enumerate(values, start=1):
            mc.cell(row=r, column=j, value=v)
    m_first, m_last = hr + 1, hr + len(models)
    style_body(mc, m_first, m_last, 11, {3: "0", 4: "0.0000", 5: "0.000", 6: "0.000", 7: "0.0",
                                          8: "0.000", 9: "0.000", 10: "0.0"})
    for r in range(m_first, m_last + 1):
        mc.cell(row=r, column=1).font = bold_font
        mc.cell(row=r, column=1).alignment = left
        mc.cell(row=r, column=2).alignment = left
    mc.conditional_formatting.add(
        f"K{m_first}:K{m_last}",
        FormulaRule(formula=[f'K{m_first}="Yes"'], fill=PatternFill("solid", start_color="C8E6C9", end_color="C8E6C9"),
                    font=Font(name=FONT, bold=True, color="1B5E20")))
    mc.conditional_formatting.add(
        f"K{m_first}:K{m_last}",
        FormulaRule(formula=[f'K{m_first}="No"'], fill=PatternFill("solid", start_color="FFCDD2", end_color="FFCDD2"),
                    font=Font(name=FONT, bold=True, color="B71C1C")))

    sig = lambda p: "significant" if p < 0.05 else "not significant"
    add_notes(mc, m_last + 2, [
        "How to read this",
        f"Model 1 (national): adjusted R\u00b2 = {m1['adj']:.2f}; the rainfall effect is {sig(m1['p'])} at 5%.",
        "Model 2 mixes counties, so part of the rainfall effect reflects differences between counties "
        "(soil, altitude, farming system).",
        f"Model 3 controls for those differences: the rainfall effect is {m3['coef'] * 100:.1f} kg/acre per 100 mm "
        f"and {sig(m3['p'])} at 5% (p = {m3['p']:.3f}).",
        "AIC (lower is better) can only be compared between Models 2 and 3, which use the same data.",
        "Caution: few counties and years, rows within a county are not independent, and rainfall is one "
        "satellite point per county.",
        "These values are calculated from the merged data files each time this script runs.",
    ])
    widths(mc, [28, 24, 14, 22, 11, 10, 22, 10, 12, 10, 16])
    mc.freeze_panes = "A5"

    Path("outputs").mkdir(exist_ok=True)
    wb.save(OUT_FILE)
    print(f"Saved {OUT_FILE}")
    for name, data, m in models:
        print(f"  {name:30s} {data:24s} n={m['n']:<3d} coef={m['coef']:.4f} p={m['p']:.3f} "
              f"adjR2={m['adj']:.3f} AIC={m['aic']:.1f}")


if __name__ == "__main__":
    main()
